"""scenic_align: turning the team workbook's scenics into collector records.

The first two collection runs have no `place_type=scenic` record, so the 12
scenics in the team workbook are the only local source for that category. These
tests pin the boundaries that make the alignment safe: the workbook's own
`pl_wb_*` namespace is kept, cross-namespace near-duplicates are flagged instead
of merged, district comes only from a source statement, and missing coordinates
stay unknown with a review task.

Workbooks are written with openpyxl; no network, no runtime fact store.
"""
import json
from pathlib import Path

import pytest

from collector import scenic_align as sa

STAMP = "2026-10-24T13:41:09+08:00"

# 表头放在第 3 行，与真实团队工作簿一致。
HEADER_ROW = ["place_id", "name", "place_type", "address", "description", "aliases",
              "source_id", "validation_status", "notes"]
COLUMNS = {"place_id": 0, "name": 1, "place_type": 2, "address": 3, "description": 4,
           "aliases": 5, "source_id": 6, "validation_status": 7, "notes": 8}


def write_workbook(path: Path, rows: list[dict]) -> Path:
    """Write a workbook whose 10_places sheet mirrors the real header layout."""
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.create_sheet(sa.SHEET)
    # 前两行是标题与说明，真实工作簿的表头从第 3 行开始。
    sheet.append(["团队数据工作簿 v1.1"])
    sheet.append(["10_places 地点清单"])
    sheet.append(HEADER_ROW)
    for row in rows:
        values = [""] * len(HEADER_ROW)
        for key, value in row.items():
            if key in COLUMNS:
                values[COLUMNS[key]] = value
        sheet.append(values)
    workbook.remove(workbook["Sheet"])
    workbook.save(str(path))
    return path


def scenic_row(**kwargs) -> dict:
    base = {"place_id": "pl_000001", "name": "清明上河园", "place_type": "scenic_spot",
            "address": "河南省开封市龙亭西路5号", "description": "国家AAAA级景区",
            "aliases": "清明上河园", "source_id": "src_000004",
            "validation_status": "candidate", "notes": "官网介绍"}
    base.update(kwargs)
    return base


def team_bundle(tmp_path: Path, places: list[dict], sources: list[dict]) -> Path:
    path = tmp_path / "team-bundle.json"
    path.write_text(json.dumps({"schema_version": "city-facts/1",
                                "records": sources + places}, ensure_ascii=False), encoding="utf-8")
    return path


def place_fact(**kwargs) -> dict:
    base = {"fact_id": "pl_wb_000001", "version": 1, "kind": "place", "name": "清明上河园",
            "place_type": "scenic", "address": "河南省开封市龙亭西路5号",
            "source_ids": ["src_teamworkbook_places", "src_wb_src_000004"],
            "validation_status": "unverified", "business_status": "active", "is_demo": False,
            "rights_status": "internal_only"}
    base.update(kwargs)
    return base


def source_fact(fact_id="src_teamworkbook_places") -> dict:
    return {"fact_id": fact_id, "version": 1, "kind": "source", "name": fact_id,
            "source_type": "other", "url": None, "locator": "workbook#10_places",
            "collected_at": "2026-10-23T22:49:00+08:00", "source_ids": [],
            "validation_status": "unverified", "business_status": "active", "is_demo": False,
            "rights_status": "internal_only"}


def build(workbook: Path, team: Path, **kwargs):
    return sa.build_records(sa.load_scenic_sources(workbook), team_bundle=team,
                            collected_at=STAMP, **kwargs)


# ------------------------------------------------------------------ reading

def test_load_scenic_sources_reads_only_scenic_rows(tmp_path):
    path = write_workbook(tmp_path / "wb.xlsx", [
        scenic_row(),
        scenic_row(place_id="pl_000002", name="开封府", place_type="museum"),
        {},  # 空行
    ])
    sources = sa.load_scenic_sources(path)
    assert [s.workbook_place_id for s in sources] == ["pl_000001"]
    assert sources[0].name == "清明上河园"
    assert sources[0].source_id == "src_000004"
    assert sources[0].workbook_status == "candidate"


def test_load_scenic_sources_skips_rows_without_id_or_name(tmp_path):
    path = write_workbook(tmp_path / "wb.xlsx", [
        scenic_row(),
        scenic_row(place_id="", name="无名景区"),
        scenic_row(place_id="pl_000003", name=""),
    ])
    assert [s.workbook_place_id for s in sa.load_scenic_sources(path)] == ["pl_000001"]


def test_load_scenic_sources_rejects_missing_sheet(tmp_path):
    import openpyxl
    path = tmp_path / "wb.xlsx"
    openpyxl.Workbook().save(str(path))
    with pytest.raises(FileNotFoundError):
        sa.load_scenic_sources(path)


def test_load_scenic_sources_rejects_headerless_sheet(tmp_path):
    import openpyxl
    path = tmp_path / "wb.xlsx"
    workbook = openpyxl.Workbook()
    workbook.create_sheet(sa.SHEET).append(["something", "else"])
    workbook.remove(workbook["Sheet"])
    workbook.save(str(path))
    with pytest.raises(ValueError):
        sa.load_scenic_sources(path)


# ------------------------------------------------------------------ district hints

def test_district_hints_strip_the_district_prefix():
    """正则容易把"龙亭区清明上河园景区"整段吃掉；城区前缀必须先剥掉。"""
    hints = sa.district_hints_from_existing(
        [{"name": "开封市常明古灯博物馆",
          "address": "开封市龙亭区清明上河园景区内上善门"}])
    assert hints["清明上河园"] == "龙亭区"
    assert hints["清明上河园景区"] == "龙亭区"


def test_district_hints_ignore_addresses_without_a_district():
    assert sa.district_hints_from_existing(
        [{"name": "开封府", "address": "包公东湖北岸"}]) == {}


def test_district_hints_ignore_short_fragments():
    assert sa.district_hints_from_existing(
        [{"name": "铁塔", "address": "龙亭区北门大街"}]) == {}


# ------------------------------------------------------------------ build_records

def test_build_records_keeps_the_workbook_namespace(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, report = build(workbook, team)
    assert [r.record_id for r in records] == ["pl_wb_000001"]
    assert records[0].place_type == "scenic"
    assert records[0].category == "scenic"
    assert records[0].source_keys == ["src_teamworkbook_places", "src_wb_src_000004"]
    assert len(report.scenics) == 1
    assert report.scenics[0]["team_bundle_present"] is True


def test_build_records_falls_back_to_the_workbook_source_when_none_declared(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(source_id="")])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, _ = build(workbook, team)
    assert records[0].source_keys == ["src_teamworkbook_places"]


def test_build_records_takes_district_from_the_workbook_address(tmp_path):
    """工作簿地址写出城区名时直接用；"龙亭西路5号"这种不含城区名的地址不算。"""
    workbook = write_workbook(tmp_path / "wb.xlsx",
                              [scenic_row(address="开封市鼓楼区包公东湖北岸")])
    team = team_bundle(tmp_path, [place_fact(address="开封市鼓楼区包公东湖北岸")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    records, report = build(workbook, team)
    assert records[0].district == "鼓楼区"
    assert report.scenics[0]["district_source"] == "工作簿地址文本"
    assert "属来源声明而非核验" not in records[0].confidence_note


def test_build_records_does_not_read_a_district_out_of_a_street_name(tmp_path):
    """"龙亭西路"不是"龙亭区"：不含城区名的地址必须留空并记核验。"""
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, report = build(workbook, team)
    assert records[0].district is None
    assert report.scenics[0]["district"] is None
    assert any("城区归属待核验" in i for i in records[0].review_items)


def test_build_records_uses_an_existing_run_district_as_a_labelled_hint(tmp_path):
    """工作簿没写城区时不猜：只采用既有 run 在同景区地址上的声明并标注。"""
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(address="开封市宋都御街北段")])
    team = team_bundle(tmp_path, [place_fact(address="开封市宋都御街北段")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    existing = [{"record_id": "place_yigutang_museum", "run_id": "run_1d70ae0fdd19",
                 "name": "开封市弋古堂文房博物馆",
                 "address": "开封市龙亭区清明上河园景区内四方院"}]
    records, report = build(workbook, team, existing_places=existing)
    assert records[0].district == "龙亭区"
    assert "既有 run 的同景区地址声明" in report.scenics[0]["district_source"]
    assert "属来源声明而非核验" in records[0].confidence_note


def test_build_records_leaves_district_null_without_any_source(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(address="包公东湖北岸")])
    team = team_bundle(tmp_path, [place_fact(address="包公东湖北岸")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    records, report = build(workbook, team)
    assert records[0].district is None
    assert any("城区归属待核验" in i for i in records[0].review_items)
    assert report.scenics[0]["district"] is None


def test_build_records_adopts_backfilled_coordinates_only_with_a_system(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    backfill = {"matched": [{"record_id": "pl_wb_000001", "longitude": 114.33517,
                             "latitude": 34.801001, "coordinate_system": "GCJ02",
                             "source": "团队工作簿 16_food_vendors", "evidence": "高德坐标"}]}
    records, report = build(workbook, team, backfill=backfill)
    assert records[0].longitude == 114.33517
    assert records[0].coordinate_system == "GCJ02"
    assert "坐标来源" in records[0].evidence_note
    assert report.coordinates_from_backfill[0]["fact_id"] == "pl_wb_000001"
    assert report.coordinates_missing == []


def test_build_records_refuses_a_coordinate_without_a_system(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    backfill = {"matched": [{"record_id": "pl_wb_000001", "longitude": 114.33517,
                             "latitude": 34.801001, "coordinate_system": "unknown",
                             "source": "x", "evidence": "x"}]}
    records, report = build(workbook, team, backfill=backfill)
    assert records[0].longitude is None
    assert records[0].coordinate_system == "unknown"
    assert any("景区坐标缺失" in i for i in records[0].review_items)
    assert len(report.coordinates_missing) == 1


def test_build_records_leaves_missing_coordinates_as_a_review_task(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, report = build(workbook, team)
    assert records[0].longitude is None and records[0].latitude is None
    assert records[0].coordinate_system == "unknown"
    assert any("景区坐标缺失" in i for i in records[0].review_items)
    assert report.coordinates_missing[0]["fact_id"] == "pl_wb_000001"


def test_build_records_records_workbook_conflict_status(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(validation_status="pending_review")])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, _ = build(workbook, team)
    assert any("pending_review" in i for i in records[0].review_items)


def test_build_records_keeps_the_evidence_quotes_from_the_workbook(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row()])
    team = team_bundle(tmp_path, [place_fact()], [source_fact(), source_fact("src_wb_src_000004")])
    records, _ = build(workbook, team)
    assert records[0].evidence_quotes["address"] == "河南省开封市龙亭西路5号"
    assert records[0].evidence_quotes["description"] == "国家AAAA级景区"


# ------------------------------------------------------------------ cross-namespace duplicates

def test_contains_match_is_conservative():
    assert sa._contains_match("山陕甘会馆", "开封市艺术博物馆（山陕甘会馆）")
    assert not sa._contains_match("清明上河园", "万岁山武侠城")
    # 短名互相包含时不报，避免"开封"把"开封府"卷进来。
    assert not sa._contains_match("开封", "开封府")


def test_cross_namespace_duplicate_is_flagged_not_merged(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(name="山陕甘会馆")])
    team = team_bundle(tmp_path, [place_fact(name="山陕甘会馆")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    existing = [{"record_id": "place_shanshangan_museum", "run_id": "run_1d70ae0fdd19",
                 "name": "开封市艺术博物馆（山陕甘会馆）",
                 "address": "开封市龙亭区徐府街"}]
    records, report = build(workbook, team, existing_places=existing)
    # 两套记录都保留，只生成核验任务。
    assert len(records) == 1
    assert len(report.cross_namespace) == 1
    flagged = report.cross_namespace[0]
    assert flagged["kept"] == ["place_shanshangan_museum", "pl_wb_000001"]
    assert flagged["existing_run"] == "run_1d70ae0fdd19"
    assert any("跨命名空间潜在重复" in i for i in records[0].review_items)


def test_cross_namespace_duplicate_within_this_batch_flags_both_sides(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [
        scenic_row(place_id="pl_000011", name="山陕甘会馆"),
        scenic_row(place_id="pl_000012", name="延庆观"),
        scenic_row(place_id="pl_000013", name="开封市艺术博物馆（山陕甘会馆）"),
    ])
    team = team_bundle(tmp_path, [], [source_fact()])
    records, report = build(workbook, team)
    assert len(records) == 3
    assert len(report.cross_namespace) == 1
    # 后写入的一条被标出，先写入的也补上同一条核验项。
    assert any("跨命名空间潜在重复" in i for i in records[0].review_items)
    assert any("跨命名空间潜在重复" in i for i in records[2].review_items)


def test_unrelated_scenics_are_not_flagged_as_duplicates(tmp_path):
    workbook = write_workbook(tmp_path / "wb.xlsx", [
        scenic_row(place_id="pl_000001", name="清明上河园"),
        scenic_row(place_id="pl_000002", name="开封府", address="开封市鼓楼区包公东湖北岸"),
    ])
    team = team_bundle(tmp_path, [], [source_fact()])
    records, report = build(workbook, team)
    assert report.cross_namespace == []
    assert not any("跨命名空间" in i for r in records for i in r.review_items)


# ------------------------------------------------------------------ CLI artefacts

def test_main_writes_structured_records_and_alignment_report(tmp_path, monkeypatch, capsys):
    workbook = write_workbook(tmp_path / "wb.xlsx", [scenic_row(), scenic_row(name="开封府")])
    team = team_bundle(tmp_path, [place_fact()], [source_fact()])
    out = tmp_path / "runs" / "scenic_align_20261024"
    monkeypatch.setattr(sa, "WORKBOOK_GLOB", "*.xlsx")
    monkeypatch.chdir(tmp_path)
    assert sa.main(["--workbook", str(workbook), "--team-bundle", str(team),
                    "--out", str(out)]) == 0
    payload = json.loads((out / "structured" / "structured.json").read_text(encoding="utf-8"))
    assert payload["schema_version"] == "kaifeng-collection-structured/1"
    assert payload["run_id"] == out.name
    assert len(payload["records"]) == 2
    report = json.loads((out / "reports" / "scenic-alignment.json").read_text(encoding="utf-8"))
    assert report["scenics_total"] == 2  # to_dict() 汇总字段
    printed = json.loads(capsys.readouterr().out)
    assert printed["status"] == "ok" and printed["scenics"] == 2


def test_main_rejects_when_no_workbook_exists(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(sa, "WORKBOOK_GLOB", "nonexistent-*.xlsx")
    monkeypatch.chdir(tmp_path)
    assert sa.main(["--out", str(tmp_path / "out")]) == 1
    assert json.loads(capsys.readouterr().err)["status"] == "rejected"


# ------------------------------------------------------------------ scenic_run driver

def _run_dir(tmp_path, records: list[dict]) -> Path:
    directory = tmp_path / "runs" / "scenic_align_20261024"
    (directory / "structured").mkdir(parents=True)
    (directory / "structured" / "structured.json").write_text(json.dumps({
        "schema_version": "kaifeng-collection-structured/1",
        "run_id": directory.name, "collected_at": STAMP, "records": records}, ensure_ascii=False),
        encoding="utf-8")
    return directory


def _scenic_record(**kwargs):
    base = {"record_id": "pl_wb_000001", "category": "scenic", "place_type": "scenic",
            "name": "清明上河园", "district": "鼓楼区", "address": "开封市鼓楼区包公东湖北岸",
            "source_keys": ["src_teamworkbook_places", "src_wb_src_000004"],
            "collected_at": STAMP, "extraction_method": "team_workbook_places_sheet",
            "coordinate_system": "unknown", "review_items": []}
    base.update(kwargs)
    return base


@pytest.fixture
def team_workbook(tmp_path):
    return write_workbook(tmp_path / "wb.xlsx", [scenic_row(address="开封市鼓楼区包公东湖北岸")])


def test_finish_writes_the_full_artefact_chain(tmp_path, db_path, monkeypatch,
                                               team_workbook):
    from collector import scenic_run as sr
    team = team_bundle(tmp_path, [place_fact(address="开封市鼓楼区包公东湖北岸")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    monkeypatch.setattr(sr, "TEAM_BUNDLE", team)
    monkeypatch.setattr(sr, "_dry_run",
                        lambda path: {"status": "dry_run", "received": 3, "changed": 3})
    run_dir = _run_dir(tmp_path, [_scenic_record()])
    summary = sr.finish(run_dir, db_path=db_path, workbook=team_workbook,
                   expected_sha256=None)
    assert summary["records"] == 1 and summary["scenics"] == 1
    assert summary["dry_run"] == {"status": "dry_run", "received": 3, "changed": 3}
    assert (run_dir / "structured" / "validation.json").exists()
    assert (run_dir / "bundle" / "city-facts-candidate.json").exists()
    assert (run_dir / "reports" / "city-facts-dry-run.json").exists()
    assert (run_dir / "reports" / "quality-report-2026-10-24.md").exists()
    assert (run_dir / "reports" / "run-summary.json").exists()


def test_finish_carries_the_referenced_team_records_into_the_bundle(tmp_path, db_path,
                                                                   monkeypatch,
                                                                   team_workbook):
    """景区记录引用的团队工作簿来源必须随包带入，否则 dry-run 引用完整性会整包拒绝。"""
    from collector import scenic_run as sr
    team = team_bundle(tmp_path, [place_fact(address="开封市鼓楼区包公东湖北岸")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    monkeypatch.setattr(sr, "TEAM_BUNDLE", team)
    monkeypatch.setattr(sr, "_dry_run", lambda path: None)
    run_dir = _run_dir(tmp_path, [_scenic_record()])
    sr.finish(run_dir, db_path=db_path, workbook=team_workbook,
                   expected_sha256=None)
    payload = json.loads((run_dir / "bundle" / "city-facts-candidate.json").read_text(encoding="utf-8"))
    ids = [r["fact_id"] for r in payload["records"]]
    assert "pl_wb_000001" in ids
    assert "src_teamworkbook_places" in ids and "src_wb_src_000004" in ids
    # 同一 id 不得重复出现。
    assert len(ids) == len(set(ids))
    carried = next(r for r in payload["records"] if r["fact_id"] == "src_wb_src_000004")
    assert carried["kind"] == "source" and carried["validation_status"] == "unverified"


def test_finish_registers_the_workbook_with_its_real_hash(tmp_path, db_path, monkeypatch,
                                                          team_workbook):
    from collector import scenic_run as sr
    from collector import ledger
    from collector.fetch import sha256_file
    team = team_bundle(tmp_path, [], [source_fact(), source_fact("src_wb_src_000004")])
    monkeypatch.setattr(sr, "TEAM_BUNDLE", team)
    monkeypatch.setattr(sr, "_dry_run", lambda path: None)
    run_dir = _run_dir(tmp_path, [_scenic_record()])
    sr.finish(run_dir, db_path=db_path, workbook=team_workbook, expected_sha256=None)
    with ledger.connect(db_path) as db:
        row = ledger.one(db, "SELECT * FROM source_registry WHERE source_key=?",
                         ("src_teamworkbook_places",))
        assert row["content_sha256"] == sha256_file(team_workbook)
        assert row["raw_path"] == str(team_workbook)
        assert row["access_status"] == "reachable"
        missing = ledger.one(db, "SELECT * FROM source_registry WHERE source_key=?",
                             ("src_wb_src_000004",))
        # 工作簿声明的证据文件不在本仓库：如实标 missing，不伪造可达。
        assert missing["access_status"] == "missing_evidence"
        assert missing["content_sha256"] is None
        run = ledger.get_run(db, run_dir.name)
        assert run["status"] == "success"


def test_finish_rejects_a_workbook_whose_hash_changed(tmp_path, db_path, monkeypatch):
    """工作簿内容变化时必须停下，不能拿旧哈希当新文件的证据。"""
    from collector import scenic_run as sr
    team = team_bundle(tmp_path, [], [source_fact()])
    monkeypatch.setattr(sr, "TEAM_BUNDLE", team)
    tampered = write_workbook(tmp_path / "tampered.xlsx", [scenic_row(name="另一个景区")])
    run_dir = _run_dir(tmp_path, [_scenic_record()])
    with pytest.raises(ValueError, match="SHA-256"):
        sr.finish(run_dir, db_path=db_path, workbook=tampered,
                  expected_sha256="0" * 64)


def test_finish_registers_review_tasks_idempotently(tmp_path, db_path, monkeypatch,
                                                    team_workbook):
    from collector import scenic_run as sr
    from collector import ledger
    team = team_bundle(tmp_path, [place_fact(address="开封市鼓楼区包公东湖北岸")],
                       [source_fact(), source_fact("src_wb_src_000004")])
    monkeypatch.setattr(sr, "TEAM_BUNDLE", team)
    monkeypatch.setattr(sr, "_dry_run", lambda path: None)
    run_dir = _run_dir(tmp_path, [_scenic_record()])
    first = sr.finish(run_dir, db_path=db_path, workbook=team_workbook,
                   expected_sha256=None)
    second = sr.finish(run_dir, db_path=db_path, workbook=team_workbook,
                   expected_sha256=None)
    with ledger.connect(db_path) as db:
        open_tasks = ledger.list_verify_tasks(db, status="open")
    assert first["review_tasks_created"] == len(open_tasks)
    assert second["review_tasks_created"] == len(open_tasks)


def test_dry_run_returns_none_when_script_missing(tmp_path, monkeypatch):
    from collector import scenic_run as sr
    monkeypatch.setattr(sr, "REPO_ROOT", tmp_path)
    assert sr._dry_run(tmp_path / "bundle.json") is None
