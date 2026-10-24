"""Coordinate backfill (Skill C engine, coordinate half).

Attaches coordinates that already exist in a *local, already-captured* source to
place records that are missing them:

* the team handoff workbook (`data/team-handoff/*.xlsx`, GCJ-02),
* the OSM POI table from the 2026-10-20 snapshot (`osm_pois.csv`, WGS84).

Nothing here calls a geocoding API and nothing invents a position. A coordinate
is only carried over when the name match is exact or a qualified strong match
(see `name_match`), and the coordinate system declared by the source is
preserved: OSM stays WGS84, the workbook stays GCJ02. No conversion happens,
because a silent WGS84/GCJ02 shift is a ~hundreds-of-metres error that would
fabricate a route. Places with no defensible match keep `null` and get a review
task instead.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
from dataclasses import dataclass, field
from pathlib import Path

from . import ledger
from .config import DEFAULT_DB_PATH, DEFAULT_RUNS_DIR, REPO_ROOT
from .extract import StructuredBundle, StructuredRecord
from .validate import normalize, normalize_name

WORKBOOK_GLOB = "data/team-handoff/*.xlsx"
OSM_POIS = Path("/Users/mac/Desktop/大挑/开封旅游数据采集_20261020/tables/osm_pois.csv")

# 头部/限定词里的同义商店后缀。只在"双方都带括号限定词"或"一方把限定词
# 并入主名"时参与比较，避免把不同分店并成一个点。
_QUAL_NOISE = ("老店", "总店", "直营店", "旗舰店", "分店", "门店")


@dataclass
class CoordinateCandidate:
    longitude: float
    latitude: float
    coordinate_system: str
    source: str            # 人类可读来源标识
    evidence: str          # 这一坐标的来源说明（许可/证据等级）
    match_kind: str        # exact | strong
    record_name: str
    source_key: str | None = None  # city-facts/1 来源 ID；命中时写入 source_keys


@dataclass
class BackfillReport:
    matched: list[dict] = field(default_factory=list)
    unmatched: list[dict] = field(default_factory=list)
    conflicts: list[dict] = field(default_factory=list)
    coordinate_sources: dict[str, int] = field(default_factory=dict)

    @property
    def matched_count(self) -> int:
        return len(self.matched)

    @property
    def unmatched_count(self) -> int:
        return len(self.unmatched)

    @property
    def conflict_count(self) -> int:
        return len(self.conflicts)

    def to_dict(self) -> dict:
        return {
            "matched": list(self.matched), "unmatched": list(self.unmatched),
            "conflicts": list(self.conflicts),
            "coordinate_sources": dict(self.coordinate_sources),
            "matched_count": len(self.matched), "unmatched_count": len(self.unmatched),
            "conflict_count": len(self.conflicts),
        }


def _split_name(name: str) -> tuple[str, str]:
    """Split a normalised name into (head, qualifier)."""
    text = normalize_name(name)
    match = re.match(r"^(.*?)（([^）]*)）$", text)
    head, qualifier = (match.group(1), match.group(2)) if match else (text, "")
    for noise in _QUAL_NOISE:
        head = head.replace(noise, "")
        qualifier = qualifier.replace(noise, "")
    return head, qualifier.replace("店", "")


def _shingles(text: str) -> set[str]:
    """Two-character shingles, used to compare qualifiers.

    Single-character overlap is too weak: "大梁路" and "宋城路" share only 路,
    which would wrongly merge two different branches into one coordinate.
    """
    if len(text) < 2:
        return {text} if text else set()
    return {text[i:i + 2] for i in range(len(text) - 1)}


def name_match(left: str, right: str) -> str | None:
    """Return 'exact', 'strong' or None for two place names.

    Exact means the normalised names are identical, including the branch
    qualifier. Strong is granted when the bracketed qualifiers share a
    two-character shingle ("寺后街" on both sides) and the shop names share a
    two-character prefix, so "第一楼（寺后街店）" matches
    "第一楼·中华老字号（寺后街总店）" and the alias form
    "邢家锅贴老店（大梁门总店）" matches "邢家锅贴大梁门总店". Two different branches
    ("黄家老店（西司店）" vs "黄家老店（宋城路店）") never match, and neither
    do plain prefix pairs such as "开封市博物馆" vs "开封市博物馆新馆" --
    those are different places and are left to human review.
    """
    if normalize_name(left) == normalize_name(right):
        return "exact"
    head_a, qual_a = _split_name(left)
    head_b, qual_b = _split_name(right)
    if (qual_a and not qual_b) or (qual_b and not qual_a):
        long_head, short_head = (head_a, head_b) if len(head_a) > len(head_b) else (head_b, head_a)
        qualifier = qual_a or qual_b
        if not short_head or not long_head.startswith(short_head):
            return None
        tail = long_head[len(short_head):]
        return "strong" if _shingles(qualifier) & _shingles(tail) else None
    if qual_a and qual_b:
        # 双方都带限定词：限定词须共享二字片段，且招牌名不是完全不同的写法
        # （"第一楼" vs "第一楼中华老字号" 是同一家店的全称/简称）。
        if not (_shingles(qual_a) & _shingles(qual_b)):
            return None
        if len(head_a) < 2 or len(head_b) < 2 or head_a[:2] != head_b[:2]:
            return None
        if head_a == head_b or head_a.startswith(head_b) or head_b.startswith(head_a):
            return "strong"
        return "strong" if _shingles(head_a) & _shingles(head_b) else None
    return None


def _crs(raw: str | None) -> str | None:
    """Map a source CRS spelling onto the city-facts/1 literal.

    EPSG:4326 is the same datum as WGS84, so it is the one alias accepted here.
    Anything else (including "未明确标注") returns None: an undeclared
    coordinate is not adopted, because guessing the datum would shift the point.
    """
    text = (raw or "").strip().upper().replace("-", "").replace("_", "").replace(" ", "")
    if text in ("GCJ02", "WGS84"):
        return text
    if text in ("EPSG4326", "EPSG:4326"):
        return "WGS84"
    return None


def load_workbook_candidates(path: Path) -> list[CoordinateCandidate]:
    """Read GCJ-02 coordinates from the team handoff workbook's vendor table."""
    import openpyxl
    candidates: list[CoordinateCandidate] = []
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        for sheet in ("16_food_vendors", "10_places"):
            if sheet not in workbook.sheetnames:
                continue
            rows = list(workbook[sheet].iter_rows(values_only=True))
            # 表头行不固定（真实工作簿前两行是标题/说明），按内容探测：
            # 同时含 name 与 lng/lat 列的第一行才是表头。
            header_row = next((i for i, row in enumerate(rows)
                               if row and "name" in [str(h).strip() if h else "" for h in row]
                               and any(str(h).strip() in ("lng", "lat") for h in row if h)), None)
            if header_row is None:
                continue
            header = [str(h).strip() if h else "" for h in rows[header_row]]
            index = {h: i for i, h in enumerate(header) if h}
            for row in rows[header_row + 1:]:
                if not row or not any(v not in (None, "") for v in row):
                    continue
                name = row[index["name"]] if "name" in index else None
                if not name:
                    continue
                lng = row[index["lng"]] if "lng" in index else None
                lat = row[index["lat"]] if "lat" in index else None
                if lng in (None, "") or lat in (None, ""):
                    continue
                system = _crs(str(row[index["crs"]] if "crs" in index and row[index["crs"]] else ""))
                if system is None:
                    continue  # 未声明坐标系的坐标不采用
                candidates.append(CoordinateCandidate(
                    longitude=float(lng), latitude=float(lat), coordinate_system=system,
                    source=f"team_workbook:{path.name}:{sheet}",
                    evidence="团队交接工作簿；高德地点页坐标，未现场核验",
                    match_kind="", record_name=str(name).strip(),
                    source_key=f"src_teamworkbook_{'food_vendors' if sheet == '16_food_vendors' else 'places'}"))
    finally:
        workbook.close()
    return candidates


def load_osm_candidates(path: Path) -> list[CoordinateCandidate]:
    """Read point coordinates from the OSM POI snapshot (WGS84, ODbL-1.0)."""
    candidates: list[CoordinateCandidate] = []
    with open(path, encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            name = (row.get("name") or "").strip()
            lng = (row.get("longitude") or "").strip()
            lat = (row.get("latitude") or "").strip()
            if not (name and lng and lat):
                continue
            system = _crs(row.get("coordinate_system"))
            if system is None:
                continue
            candidates.append(CoordinateCandidate(
                longitude=float(lng), latitude=float(lat), coordinate_system=system,
                source=f"osm:{path.name}:{row.get('osm_type')}/{row.get('osm_id')}",
                evidence="OpenStreetMap 快照；osm_mapped_not_field_verified；许可 ODbL-1.0（署名要求见 bundle rights 说明）",
                match_kind="", record_name=name, source_key="src_osm_pois"))
    return candidates


def load_candidates(*, workbook_paths: list[Path] | None = None,
                    osm_path: Path | None = None) -> list[CoordinateCandidate]:
    """Collect coordinate candidates from every local source that is present."""
    out: list[CoordinateCandidate] = []
    for path in (workbook_paths if workbook_paths is not None
                 else sorted(REPO_ROOT.glob(WORKBOOK_GLOB))):
        if path.exists():
            out.extend(load_workbook_candidates(path))
    table = osm_path if osm_path is not None else OSM_POIS
    if table.exists():
        out.extend(load_osm_candidates(table))
    return out


def backfill(records: list[StructuredRecord], candidates: list[CoordinateCandidate],
             *, place_categories: set[str]) -> tuple[list[dict], BackfillReport]:
    """Attach coordinates to records. Returns (updated records, report).

    A record is only updated when exactly one defensible candidate matches. Two
    different candidates for the same record are a conflict: both values are
    reported and the record keeps `null`, because silently choosing one would
    fabricate a position.
    """
    report = BackfillReport()
    updated: list[StructuredRecord] = []
    by_name: dict[str, list[CoordinateCandidate]] = {}
    for candidate in candidates:
        by_name.setdefault(normalize_name(candidate.record_name), []).append(candidate)

    for record in records:
        if record.category not in place_categories:
            updated.append(record)
            continue
        if record.longitude is not None and record.latitude is not None:
            updated.append(record)
            continue
        hits: list[tuple[CoordinateCandidate, str]] = []
        for key, group in by_name.items():
            kind = name_match(record.name or "", key)
            if kind is None:
                continue
            for candidate in group:
                hits.append((candidate, kind))
        if not hits:
            report.unmatched.append({
                "record_id": record.record_id, "name": record.name,
                "category": record.category, "reason": "本地来源无同名/同名分店坐标候选"})
            updated.append(record)
            continue
        # 未声明坐标系的候选一律不采用：调用方可能手工构造候选，
        # 这里再兜底一次，避免 unknown/local_demo 混进 place 层。
        hits = [(c, k) for c, k in hits if c.coordinate_system in ("WGS84", "GCJ02")]
        if not hits:
            report.unmatched.append({
                "record_id": record.record_id, "name": record.name,
                "category": record.category, "reason": "坐标候选未声明坐标系，不采用"})
            updated.append(record)
            continue
        distinct = {(round(c.longitude, 6), round(c.latitude, 6), c.coordinate_system)
                    for c, _ in hits}
        if len(distinct) > 1:
            report.conflicts.append({
                "record_id": record.record_id, "name": record.name,
                "values": [{"source": c.source, "longitude": c.longitude,
                            "latitude": c.latitude, "coordinate_system": c.coordinate_system,
                            "match_kind": kind} for c, kind in hits],
                "resolution": "候选坐标互相矛盾，保留 null，交由人工核验"})
            updated.append(record)
            continue
        candidate, kind = hits[0]
        record.longitude = candidate.longitude
        record.latitude = candidate.latitude
        record.coordinate_system = candidate.coordinate_system
        if candidate.source_key and candidate.source_key not in record.source_keys:
            record.source_keys = list(record.source_keys) + [candidate.source_key]
        record.evidence_note = (f"{record.evidence_note}｜坐标来源：{candidate.source}"
                                if record.evidence_note else f"坐标来源：{candidate.source}")
        report.matched.append({
            "record_id": record.record_id, "name": record.name, "category": record.category,
            "longitude": candidate.longitude, "latitude": candidate.latitude,
            "coordinate_system": candidate.coordinate_system, "source": candidate.source,
            "match_kind": kind, "evidence": candidate.evidence})
        report.coordinate_sources[candidate.coordinate_system] = \
            report.coordinate_sources.get(candidate.coordinate_system, 0) + 1
        updated.append(record)
    return updated, report


PLACE_CATEGORIES = {"scenic", "museum", "street_block", "park", "night_market",
                    "restaurant", "transport", "rest", "entrance"}


def build_backfilled_bundle(source_run_ids: list[str], candidates: list[CoordinateCandidate],
                            *, runs_dir: Path) -> tuple[StructuredBundle, BackfillReport]:
    """Merge earlier runs' structured records and attach coordinates.

    Records keep their original `record_id`, so the produced bundle is a
    coordinate-only delta against the earlier runs rather than a new set of
    places. Later runs win on record_id collisions.
    """
    merged: dict[str, StructuredRecord] = {}
    for run_id in source_run_ids:
        path = Path(runs_dir) / run_id / "structured" / "structured.json"
        if not path.exists():
            raise FileNotFoundError(f"缺少结构化产物：{path}")
        data = json.loads(path.read_text(encoding="utf-8"))
        for item in data.get("records", []):
            merged[item["record_id"]] = StructuredRecord.from_dict(item)
    records, report = backfill(list(merged.values()), candidates,
                               place_categories=PLACE_CATEGORIES)
    bundle = StructuredBundle(run_id=f"coord_backfill_{'+'.join(source_run_ids)}",
                              collected_at=ledger.now_iso())
    bundle.records = records
    return bundle, report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-ids", nargs="+", required=True,
                        help="要回填坐标的历史 run_id（按顺序合并）")
    parser.add_argument("--database", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--runs-dir", type=Path, default=DEFAULT_RUNS_DIR)
    parser.add_argument("--osm-pois", type=Path, default=OSM_POIS)
    parser.add_argument("--out", type=Path, default=None,
                        help="输出目录；默认写到 runs-dir/coord_backfill_<run_ids>")
    parser.add_argument("--dry-run", action="store_true",
                        help="只打印匹配报告，不落盘")
    args = parser.parse_args(argv)

    candidates = load_candidates(osm_path=args.osm_pois)
    if not candidates:
        print(json.dumps({"status": "rejected",
                          "error": "没有可用坐标候选（工作簿与 OSM 表都不存在或无可采用坐标）"},
                         ensure_ascii=False), file=sys.stderr)
        return 1
    bundle, report = build_backfilled_bundle(args.run_ids, candidates, runs_dir=args.runs_dir)
    if args.dry_run:
        print(json.dumps(report.to_dict(), ensure_ascii=False, indent=2))
        return 0

    out_dir = args.out or (Path(args.runs_dir) / f"coord_backfill_{'_'.join(args.run_ids)}")
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "structured").mkdir(exist_ok=True)
    (out_dir / "bundle").mkdir(exist_ok=True)
    (out_dir / "reports").mkdir(exist_ok=True)
    bundle.write(out_dir / "structured" / "structured.json")
    (out_dir / "reports" / "coordinate-backfill.json").write_text(
        json.dumps(report.to_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "ok", "run_dir": str(out_dir),
                      "matched": len(report.matched), "unmatched": len(report.unmatched),
                      "conflicts": len(report.conflicts),
                      "coordinate_sources": report.coordinate_sources},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
